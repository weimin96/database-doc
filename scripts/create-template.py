"""从视觉母版提取样式，构建不含旧项目内容的受控模板。仅开发时运行。"""
from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

root = Path(__file__).resolve().parent.parent
source = Document(root / '中规院基于AI的城市设计助手数据库设计文档v0.1.docx')
doc = Document()
doc.styles.element.clear()
for style in source.styles.element:
    doc.styles.element.append(deepcopy(style))
for section in doc.sections:
    original = source.sections[-1]
    for attribute in ('page_width', 'page_height', 'top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
        setattr(section, attribute, getattr(original, attribute))
for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
    style = doc.styles[name]
    style.font.color.rgb = RGBColor(0, 0, 0)
    for tag in list(style.element.xpath('.//w:numPr')):
        tag.getparent().remove(tag)
    for color in style.element.xpath('.//w:color'):
        for attribute in ('themeColor', 'themeTint', 'themeShade'):
            color.attrib.pop(qn('w:' + attribute), None)
for style in doc.styles:
    if style.type in (1, 2):
        style.font.name = 'Songti SC'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Songti SC')
        for fonts in style.element.xpath('.//w:rFonts'):
            for key in list(fonts.attrib):
                if 'Theme' in key:
                    del fonts.attrib[key]
for name in ('Heading 1', 'Heading 2', 'Heading 3'):
    doc.styles[name].paragraph_format.page_break_before = False
    doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles[name].paragraph_format.space_before = Pt(14)
    doc.styles[name].paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(26)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].font.name = 'Songti SC'
doc.styles['Normal'].element.rPr.rFonts.set(qn('w:eastAsia'), 'Songti SC')
doc.add_paragraph('{projectName}', 'Title')
doc.add_paragraph('数据库设计说明书', 'Title')
for text in ('建设单位：{organization}', '文档版本：{version}', '编制日期：{date}'):
    doc.add_paragraph(text)
doc.add_page_break()
doc.add_heading('项目概述', 1)
doc.add_paragraph('{overview}')
doc.add_heading('数据表清单', 1)

def table(prototype, headings, values, widths):
    element = deepcopy(prototype._tbl)
    # 仅复用表格属性与空白网格；不保留历史内容及绑定关系。
    for row in list(element.findall(qn('w:tr'))):
        element.remove(row)
    doc._body._body.insert(-1, element)
    result = doc.tables[-1]
    result.autofit = False
    for col, width in zip(result.columns, widths):
        col.width = Cm(width)
    for row_values in (headings, values):
        row = result.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        cells = row.cells
        for cell, value, width in zip(cells, row_values, widths):
            cell.text = value
            cell.width = Cm(width)
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                margin = OxmlElement('w:' + side)
                margin.set(qn('w:w'), '50')
                margin.set(qn('w:type'), 'dxa')
                margins.append(margin)
            cell._tc.get_or_add_tcPr().append(margins)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.left_indent = Pt(0)
            paragraph.paragraph_format.right_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            indent = paragraph._p.get_or_add_pPr().get_or_add_ind()
            for key in ('firstLineChars', 'leftChars', 'rightChars', 'hangingChars'):
                indent.set(qn('w:' + key), '0')
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)
                run.font.name = 'Songti SC'
                run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Songti SC')
    header = result.rows[0]
    header._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for cell in header.cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True
        cell.paragraphs[0].paragraph_format.keep_with_next = True
    # 固定为可用版心宽度，去掉母版的负缩进和百分比溢出。
    properties = result._tbl.tblPr
    for item in list(properties):
        if item.tag in (qn('w:tblW'), qn('w:tblInd')):
            properties.remove(item)
    return result

table(source.tables[1], ['#', '数据表', '名称', '备注说明'],
      ['{#tables}{index}', '{qualifiedName}', '{displayName}', '{/tables}'], [1, 5, 5, 4])
doc.add_heading('数据表设计', 1)
doc.add_paragraph('{#tables}')
doc.add_heading('{qualifiedName}', 2)
description = doc.add_paragraph('{displayName}')
description.paragraph_format.keep_with_next = True
table(source.tables[2], ['#', '字段', '名称', '数据类型', '主键', '非空', '默认值', '备注说明'],
      ['{#columns}{index}', '{name}', '{displayName}', '{dataType}', '{primaryKey}', '{notNull}', '{defaultValue}', '{remarks}{/columns}'],
      [0.8, 2.4, 3.4, 2.4, 0.9, 0.9, 2.2, 2])
doc.add_paragraph('{/tables}')
doc.core_properties.title = '数据库设计说明书'
doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''
doc.core_properties.comments = ''
output = root / 'templates/database-design.template.docx'
doc.save(output)
print(output)
